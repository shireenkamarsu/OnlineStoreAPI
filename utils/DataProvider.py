import json
import csv
from openpyxl import load_workbook


# JSON Data Provider
def read_json_data(filepath):
    with open(filepath, "r") as f:
        data_list = json.load(f)
    return [(item,) for item in data_list]


# CSV Data Provider
def read_csv_data(filepath):
    data_list = []
    with open(filepath, "r") as f:
        reader = csv.reader(f)  # Create a csv reader object to read each low line by line in csv file
        next(reader)
        for row in reader:
            data_list.append(tuple(row))
    return data_list

# Excel Data Provider
def read_excel_data(filepath,sheetname):
    wb = load_workbook(filepath) # Load the Excel workbook.
    sheet = wb[sheetname] # Access the specified sheet by name.
    data = [] # Empty list to store data rows
    for row in sheet.iter_rows(min_row=2,values_only=True): #Loop through all rows from 2nd row onwards (to skip the header)
        data.append(row) # Add each row (as a tuple) to the list
    return data # Return list of tuples
